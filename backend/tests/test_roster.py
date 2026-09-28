"""Roster compatibility and validation, using fictional students only."""

import io
import zipfile
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Font

from event_api.errors import ApiError
from event_api.routers.roster import HEADERS, REGISTER_HEADERS, parse_roster


def workbook_bytes(workbook: Workbook) -> bytes:
    target = io.BytesIO()
    workbook.save(target)
    workbook.close()
    return target.getvalue()


@pytest.mark.parametrize("title", [True, False])
def test_register_accepts_header_wrapping_and_empty_styled_columns(title: bool) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    if title:
        sheet.append(["Реестр контингента"])
        sheet.merge_cells("A1:G1")
    sheet.append([header.upper().replace(" ", "\n") for header in REGISTER_HEADERS])
    sheet.append(
        [
            "Тестов Тест Тестович",
            "Обучается",
            "ТЕСТ-1",
            "Площадка",
            1,
            "Программа",
            "00.00.00",
        ]
    )
    sheet["J20"].font = Font(bold=True)
    parsed = parse_roster(workbook_bytes(book))
    assert len(parsed) == 1
    assert parsed[0]["last_name"] == "Тестов"
    assert parsed[0]["study_group"] == "ТЕСТ-1"


def test_simple_roster_ignores_formatting_beyond_data() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(list(HEADERS))
    sheet.append(["Тестов", "Тест", None, "ТЕСТ-1"])
    sheet["G5"].font = Font(bold=True)
    assert len(parse_roster(workbook_bytes(book))) == 1


@pytest.mark.parametrize("coordinate", ["D1048576", "XFD3", "Z50000"])
def test_sparse_grid_is_rejected_before_workbook_materialization(
    coordinate, monkeypatch
):
    book = Workbook()
    sheet = book.active
    sheet.append(list(HEADERS))
    sheet.append(["Тестов", "Тест", None, "ТЕСТ-1"])
    sheet[coordinate].font = Font(bold=True)
    source = workbook_bytes(book)

    def unexpected_load(*args, **kwargs):
        pytest.fail("Oversized grid reached openpyxl")

    monkeypatch.setattr("event_api.routers.roster.load_workbook", unexpected_load)
    with pytest.raises(ApiError) as failure:
        parse_roster(source)
    assert failure.value.details == {"reason": "GRID_TOO_LARGE"}


def test_oversized_merge_is_rejected_before_materialization(monkeypatch):
    book = Workbook()
    book.active.append(list(HEADERS))
    book.active.append(["Тестов", "Тест", None, "ТЕСТ-1"])
    source = workbook_bytes(book)
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(source)) as original,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in original.infolist():
            value = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                value = value.replace(
                    b"</worksheet>",
                    b'<mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells></worksheet>',
                )
            target.writestr(info, value)
    monkeypatch.setattr(
        "event_api.routers.roster.load_workbook",
        lambda *args, **kwargs: pytest.fail("Large merge reached openpyxl"),
    )
    with pytest.raises(ApiError) as failure:
        parse_roster(output.getvalue())
    assert failure.value.details == {"reason": "GRID_TOO_LARGE"}


@pytest.mark.parametrize(
    ("values", "reason", "column"),
    [
        (["Тестов", "", None, "ТЕСТ-1"], "REQUIRED_VALUE", 2),
        (["Тестов", "=1+1", None, "ТЕСТ-1"], "FORMULA", 2),
        (["Тестов", "Тест", None, "ТЕСТ-1", "Лишнее"], "EXTRA_COLUMNS", 5),
    ],
)
def test_row_errors_identify_location_without_personal_data(
    values: list[str | None], reason: str, column: int
) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(list(HEADERS))
    sheet.append(values)
    with pytest.raises(ApiError) as failure:
        parse_roster(workbook_bytes(book))
    assert failure.value.details == {"reason": reason, "row": 2, "column": column}
    assert "Тестов" not in failure.value.message


def test_duplicate_rows_and_database_length_limits() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(list(HEADERS))
    sheet.append(["Тестов", "Тест", None, "ТЕСТ-1"])
    sheet.append([None])
    sheet.append([" тестов ", "Тест", None, "ТЕСТ-1"])
    with pytest.raises(ApiError) as duplicate:
        parse_roster(workbook_bytes(book))
    assert duplicate.value.details == {
        "reason": "DUPLICATE_STUDENT",
        "row": 4,
        "firstRow": 2,
    }
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(list(HEADERS))
    sheet.append(["Тестов", "Тест", None, "x" * 101])
    with pytest.raises(ApiError) as long_value:
        parse_roster(workbook_bytes(book))
    assert long_value.value.details == {
        "reason": "VALUE_TOO_LONG",
        "row": 2,
        "column": 4,
        "maxLength": 100,
    }


@pytest.mark.parametrize("case", ["empty", "sheets", "merged", "corrupt", "oversize"])
def test_invalid_workbooks_remain_rejected(case: str) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(list(HEADERS))
    expected = "EMPTY_ROSTER"
    if case == "sheets":
        book.create_sheet("Другой лист")
        expected = "SHEET_COUNT"
    if case == "merged":
        sheet.merge_cells("A2:B2")
        expected = "MERGED_CELLS"
    source = workbook_bytes(book)
    if case == "corrupt":
        source = b"PK not a workbook"
        expected = "UNSAFE_ARCHIVE"
    if case == "oversize":
        source = b"PK" + b"x" * (5 * 1024 * 1024)
        expected = "FILE_TOO_LARGE"
    with pytest.raises(ApiError) as failure:
        parse_roster(source)
    assert failure.value.details == {"reason": expected}


@pytest.mark.parametrize("mime", ["application/octet-stream", ""])
def test_formatted_roster_preview_and_commit_with_generic_file_type(
    client: TestClient, mime: str
) -> None:
    client.cookies.clear()
    origin = {"Origin": "http://localhost:5173"}
    login = client.post(
        "/auth/login",
        headers=origin,
        json={"email": "admin@example.com", "password": "correct horse battery"},
    )
    assert login.status_code == 200
    headers = {**origin, "X-CSRF-Token": login.json()["csrfToken"]}
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Реестр контингента"])
    sheet.merge_cells("A1:G1")
    sheet.append([header.replace(" ", "\n") for header in REGISTER_HEADERS])
    sheet.append(
        [
            "Проверочный Студент Тестович",
            "Обучается",
            f"ТЕСТ-{uuid4().hex[:8]}",
            "Тестовая площадка",
            1,
            "Программа",
            "00.00.00",
        ]
    )
    sheet["J20"].font = Font(bold=True)
    file = {"file": ("roster.xlsx", workbook_bytes(book), mime)}
    assert (
        client.post(
            "/admin/activity/roster/preview", headers=origin, files=file
        ).status_code
        == 403
    )
    preview = client.post("/admin/activity/roster/preview", headers=headers, files=file)
    assert preview.status_code == 200, preview.text
    assert preview.json()["students"] == 1
    changed = client.post(
        "/admin/activity/roster/import",
        headers=headers,
        files=file,
        data={"fileHash": "0" * 64},
    )
    assert changed.status_code == 409
    imported = client.post(
        "/admin/activity/roster/import",
        headers=headers,
        files=file,
        data={"fileHash": preview.json()["fileHash"]},
    )
    assert imported.status_code == 201, imported.text
    assert imported.json() == {"created": 1, "skipped": 0}
    repeated = client.post(
        "/admin/activity/roster/preview", headers=headers, files=file
    )
    assert repeated.status_code == 409
    assert repeated.json()["error"]["code"] == "ROSTER_STUDENT_EXISTS"
    safe_preview = client.post(
        "/admin/activity/roster/preview",
        headers=headers,
        files=file,
        data={"mode": "SKIP_EXACT"},
    )
    assert safe_preview.status_code == 200, safe_preview.text
    assert safe_preview.json()["students"] == 0
    assert safe_preview.json()["skipped"] == 1
    safe_import = client.post(
        "/admin/activity/roster/import",
        headers=headers,
        files=file,
        data={"mode": "SKIP_EXACT", "fileHash": safe_preview.json()["fileHash"]},
    )
    assert safe_import.status_code == 201, safe_import.text
    assert safe_import.json() == {"created": 0, "skipped": 1}
    from sqlalchemy import text

    parsed = parse_roster(file["file"][1])[0]
    with client.app.state.database.transaction() as connection:
        connection.execute(
            text(
                "UPDATE student_roster_members m JOIN persons p ON p.id=m.person_id SET m.campus_address='Исправлено вручную' WHERE p.study_group=:group"
            ),
            {"group": parsed["study_group"]},
        )
    conflict = client.post(
        "/admin/activity/roster/preview",
        headers=headers,
        files=file,
        data={"mode": "SKIP_EXACT"},
    )
    assert conflict.status_code == 409, conflict.text
    assert conflict.json()["error"]["code"] == "ROSTER_STUDENT_CONFLICT"


def test_preview_exposes_safe_row_diagnostics(client: TestClient) -> None:
    origin = {"Origin": "http://localhost:5173"}
    login = client.post(
        "/auth/login",
        headers=origin,
        json={"email": "admin@example.com", "password": "correct horse battery"},
    )
    headers = {**origin, "X-CSRF-Token": login.json()["csrfToken"]}
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(list(HEADERS))
    sheet.append(["НеПубликоватьЭтоИмя", "Тест", None, None])
    response = client.post(
        "/admin/activity/roster/preview",
        headers=headers,
        files={
            "file": ("roster.xlsx", workbook_bytes(book), "application/octet-stream")
        },
    )
    assert response.status_code == 400
    assert response.json()["error"]["details"] == {
        "reason": "REQUIRED_VALUE",
        "row": 2,
        "column": 4,
    }
    assert "НеПубликоватьЭтоИмя" not in response.text
