"""Preview and import the student roster without publishing profiles."""

from __future__ import annotations

import hashlib
import io
from typing import Annotated, Any

from fastapi import APIRouter, Depends, File, Form, UploadFile
from openpyxl import load_workbook

from ..database import Database, execute, row
from ..dependencies import Staff, csrf_super_admin, database
from ..errors import ApiError
from ..registration_service import create_person
from ..service_utils import audit
from .excel import MAX_FILE, MAX_ROWS, XLSX_MIME, _validate_xlsx_archive

router = APIRouter(prefix="/admin/activity/roster", tags=["roster"])
HEADERS = ("Фамилия", "Имя", "Отчество", "Группа")


def parse_roster(source: bytes) -> list[dict[str, str | None]]:
    if not source or len(source) > MAX_FILE or not source.startswith(b"PK"):
        raise ApiError(
            400, "INVALID_ROSTER_FILE", "A valid XLSX up to 5 MiB is required"
        )
    _validate_xlsx_archive(source)
    try:
        workbook = load_workbook(io.BytesIO(source), read_only=False, data_only=False)
    except Exception as error:
        raise ApiError(400, "INVALID_ROSTER_FILE", "Cannot read XLSX") from error
    try:
        if len(workbook.worksheets) != 1 or workbook.worksheets[0].merged_cells.ranges:
            raise ApiError(400, "INVALID_ROSTER_FILE", "One unmerged sheet is required")
        sheet = workbook.worksheets[0]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        if headers != list(HEADERS):
            raise ApiError(
                400, "INVALID_ROSTER_HEADERS", "Use: Фамилия, Имя, Отчество, Группа"
            )
        result: list[dict[str, str | None]] = []
        seen: set[tuple[str, ...]] = set()
        for cells in sheet.iter_rows(min_row=2):
            if all(cell.value in (None, "") for cell in cells):
                continue
            if len(result) >= MAX_ROWS or len(cells) != 4:
                raise ApiError(400, "INVALID_ROSTER_FILE", "Too many rows or columns")
            if any(cell.data_type == "f" for cell in cells):
                raise ApiError(400, "INVALID_ROSTER_FILE", "Formulas are not allowed")
            values = [str(cell.value or "").strip() for cell in cells]
            if any(len(value) > 120 for value in values) or not all(
                values[index] for index in (0, 1, 3)
            ):
                raise ApiError(
                    400, "INVALID_ROSTER_ROW", "Surname, name and group are required"
                )
            key = tuple(value.casefold() for value in values)
            if key in seen:
                raise ApiError(
                    409, "DUPLICATE_ROSTER_STUDENT", "Duplicate name and group"
                )
            seen.add(key)
            result.append(
                {
                    "last_name": values[0],
                    "first_name": values[1],
                    "middle_name": values[2] or None,
                    "study_group": values[3],
                }
            )
        if not result:
            raise ApiError(400, "INVALID_ROSTER_FILE", "No students in workbook")
        return result
    finally:
        workbook.close()


def assert_new_students(
    connection: Any, tenant_id: str, students: list[dict[str, str | None]]
) -> None:
    for student in students:
        found = row(
            connection,
            """SELECT p.id FROM student_roster_members member JOIN persons p ON p.id=member.person_id
            WHERE p.tenant_id=:tenant AND p.merged_into_id IS NULL
              AND p.last_name=:last AND p.first_name=:first
              AND (p.middle_name <=> :middle) AND p.study_group=:group LIMIT 1""",
            {
                "tenant": tenant_id,
                "last": student["last_name"],
                "first": student["first_name"],
                "middle": student["middle_name"],
                "group": student["study_group"],
            },
        )
        if found:
            raise ApiError(
                409, "ROSTER_STUDENT_EXISTS", "A student is already in the roster"
            )


async def upload(file: UploadFile) -> bytes:
    if (
        not (file.filename or "").lower().endswith(".xlsx")
        or file.content_type != XLSX_MIME
    ):
        raise ApiError(400, "INVALID_ROSTER_FILE", "Only XLSX is accepted")
    return await file.read(MAX_FILE + 1)


@router.post("/preview")
async def preview(
    file: Annotated[UploadFile, File()],
    staff: Annotated[Staff, Depends(csrf_super_admin)],
    db: Annotated[Database, Depends(database)],
) -> dict[str, Any]:
    source = await upload(file)
    students = parse_roster(source)
    with db.connect() as connection:
        assert_new_students(connection, staff.tenant_id, students)
    return {"fileHash": hashlib.sha256(source).hexdigest(), "students": len(students)}


@router.post("/import", status_code=201)
async def import_roster(
    file: Annotated[UploadFile, File()],
    file_hash: Annotated[str, Form(alias="fileHash")],
    staff: Annotated[Staff, Depends(csrf_super_admin)],
    db: Annotated[Database, Depends(database)],
) -> dict[str, int]:
    source = await upload(file)
    if hashlib.sha256(source).hexdigest() != file_hash:
        raise ApiError(409, "ROSTER_FILE_CHANGED", "Workbook changed after preview")
    students = parse_roster(source)
    with db.transaction() as connection:
        organization = row(
            connection,
            "SELECT name FROM organizations WHERE id=:id AND tenant_id=:tenant FOR UPDATE",
            {"id": staff.organization_id, "tenant": staff.tenant_id},
        )
        if not organization:
            raise ApiError(404, "ORGANIZATION_NOT_FOUND", "Organization not found")
        assert_new_students(connection, staff.tenant_id, students)
        for student in students:
            person_id = create_person(
                connection,
                {
                    **student,
                    "birth_date": None,
                    "email": None,
                    "phone": None,
                    "person_type": "KAIT_STUDENT",
                    "organization": organization["name"],
                },
                staff.tenant_id,
            )
            execute(
                connection,
                "INSERT INTO student_roster_members(person_id,created_at,created_by) VALUES (:id,UTC_TIMESTAMP(3),:actor)",
                {"id": person_id, "actor": staff.id},
            )
            execute(
                connection,
                """INSERT INTO student_profiles(id,person_id,visibility,created_at,updated_at)
                VALUES (UUID(),:person,'PRIVATE',UTC_TIMESTAMP(3),UTC_TIMESTAMP(3))""",
                {"person": person_id},
            )
        audit(
            connection,
            staff.id,
            "STUDENT_ROSTER_IMPORTED",
            "Organization",
            staff.organization_id,
            {"students": len(students)},
        )
    return {"created": len(students)}
